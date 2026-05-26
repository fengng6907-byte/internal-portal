"""
Elitez FMCG Portal — FastAPI application.

Single entry-point for all API routes.  Deployed as one Vercel Python
serverless function; routed via the `rewrites` rule in vercel.json.

Local development:
    uvicorn api.index:app --reload --port 5328
"""

from __future__ import annotations

import asyncio
import io
from datetime import datetime, timezone

import feedparser
import openpyxl
from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from pydantic import BaseModel

from database import (
    compute_url_hash,
    get_kol_by_id,
    insert_launch,
    read_kol_intel,
    read_launches,
    upsert_launch,
)

app = FastAPI(docs_url="/api/docs", openapi_url="/api/openapi.json")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# ---------------------------------------------------------------------------
# Brand → distributor map
# ---------------------------------------------------------------------------

DISTRIBUTOR_MAP: dict[str, str] = {
    "Nestlé":      "DKSH",
    "Unilever":    "Auric Pacific",
    "P&G":         "Direct",
    "Oatbedient":  "Direct",
    "Meiji":       "Auric Pacific",
    "Mondelez":    "DKSH",
    "Kellogg's":   "DKSH",
    "Mars":        "Direct",
    "Danone":      "DKSH",
    "Abbott":      "Zuellig Pharma",
}

# ---------------------------------------------------------------------------
# Retail channel detector
# ---------------------------------------------------------------------------

_RETAIL_CHANNEL_SIGNALS: list[tuple[str, list[str]]] = [
    ("FairPrice",    ["fairprice", "ntuc", "ntuc fairprice"]),
    ("Cold Storage", ["cold storage"]),
    ("Sheng Siong",  ["sheng siong"]),
    ("Giant",        ["giant"]),
    ("Jaya Grocer",  ["jaya grocer"]),
    ("Guardian MY",  ["guardian my", "guardian malaysia"]),
    ("Aeon",         ["aeon malaysia", "aeon"]),
    ("Lotus's",      ["lotus's"]),
    ("Mydin",        ["mydin"]),
]


def _detect_retail_channel(text: str) -> str:
    t = text.lower()
    for channel, signals in _RETAIL_CHANNEL_SIGNALS:
        if any(s in t for s in signals):
            return channel
    return ""


# ---------------------------------------------------------------------------
# Festive tagger
# ---------------------------------------------------------------------------

_FESTIVE_MAP: dict[str, list[str]] = {
    "CNY":        ["chinese new year", "cny", "lunar new year", "ang bao", "mandarin orange", "yu sheng"],
    "Hari Raya":  ["hari raya", "raya", "aidilfitri", "ramadan", "ketupat", "kuih"],
    "Mid-Autumn": ["mid-autumn", "mooncake", "mid autumn", "lantern festival", "zhong qiu"],
}


def _tag_festive(text: str) -> str:
    t = text.lower()
    for tag, keywords in _FESTIVE_MAP.items():
        if any(kw in t for kw in keywords):
            return tag
    return "None"


# ---------------------------------------------------------------------------
# Market detector — SG vs MY signal density
# ---------------------------------------------------------------------------

_MY_SIGNALS = {
    "jaya grocer", "guardian my", "aeon malaysia", "lotus's", "malaysia",
    "kuala lumpur", " kl ", "penang", "guardian malaysia", "mydin",
}
_SG_SIGNALS = {
    "fairprice", "ntuc", "cold storage", "sheng siong", "giant",
    "singapore", " sg ", "orchard", "jewel",
}


def _detect_market(text: str) -> str:
    t = text.lower()
    my_score = sum(1 for s in _MY_SIGNALS if s in t)
    sg_score = sum(1 for s in _SG_SIGNALS if s in t)
    return "Malaysia" if my_score > sg_score else "Singapore"


def _detect_distributor(text: str) -> str:
    t = text.lower()
    for brand, dist in DISTRIBUTOR_MAP.items():
        if brand.lower() in t:
            return dist
    return "Unknown"


def _detect_brand(text: str) -> str:
    t = text.lower()
    for brand in DISTRIBUTOR_MAP:
        if brand.lower() in t:
            return brand
    return "Unknown"


# ---------------------------------------------------------------------------
# RSS scraper config
# ---------------------------------------------------------------------------

_FMCG_KEYWORDS = [
    "launch", "exclusive flavor", "exclusive flavour",
    "FairPrice", "NTUC", "Cold Storage", "Sheng Siong", "Giant",
    "new product", "new SKU", "limited edition", "product launch",
    "Singapore FMCG", "FMCG Singapore",
    "Jaya Grocer", "Guardian MY", "Aeon Malaysia", "Lotus's Malaysia",
    "Mydin", "Malaysia FMCG", "FMCG Malaysia", "KL launch",
]

_RSS_FEEDS = [
    "https://www.campaignbriefasia.com/feed/",
    "https://marketing-interactive.com/feed/",
    "https://www.retailnews.asia/feed/",
    "https://www.marketing.com.my/feed/",
    "https://brandingasia.com/feed/",
]


def _is_fmcg_relevant(title: str, summary: str) -> bool:
    text = (title + " " + summary).lower()
    return any(kw.lower() in text for kw in _FMCG_KEYWORDS)


# ---------------------------------------------------------------------------
# Request models
# ---------------------------------------------------------------------------


class ConvertRequest(BaseModel):
    kol_id: str


# ---------------------------------------------------------------------------
# Routes
# ---------------------------------------------------------------------------


@app.get("/api/launches")
async def get_launches() -> dict:
    """Return all records from the Supabase launches table."""
    try:
        data = await asyncio.to_thread(read_launches)
        return {"data": data}
    except Exception as exc:
        raise HTTPException(status_code=500, detail=str(exc))


@app.get("/api/kol-signals")
async def get_kol_signals() -> dict:
    """Return all records from the Supabase kol_intel table."""
    try:
        data = await asyncio.to_thread(read_kol_intel)
        return {"data": data}
    except Exception as exc:
        raise HTTPException(status_code=500, detail=str(exc))


@app.post("/api/scraper")
async def trigger_scraper() -> dict:
    """Fetch configured RSS feeds and persist new FMCG signals to the launches table."""
    try:
        result = await asyncio.to_thread(_run_scrape)
        return result
    except Exception as exc:
        raise HTTPException(status_code=500, detail=str(exc))


@app.post("/api/convert", status_code=201)
async def convert_to_lead(req: ConvertRequest) -> dict:
    """Convert a kol_intel record into a launches lead row in Supabase."""
    kol = await asyncio.to_thread(get_kol_by_id, req.kol_id)
    if not kol:
        raise HTTPException(status_code=404, detail=f"KOL ID '{req.kol_id}' not found")

    keywords = kol.get("detected_keywords", "")
    combined = f"{kol.get('creator_display_name', '')} {keywords}"

    row: dict = {
        "id":                   f"lead_{kol['id']}",
        "brand_name":           _detect_brand(combined),
        "product_name":         f"[KOL Signal] {kol.get('creator_display_name', 'Unknown')}",
        "market":               kol.get("market", "Singapore"),
        "retail_channel":       _detect_retail_channel(combined),
        "distributor_point":    _detect_distributor(combined),
        "festive_tag":          _tag_festive(combined),
        "product_photo_url":    kol.get("product_photo_url"),
        "creator_display_name": kol.get("creator_display_name"),
        "pitch_status":         "KOL Lead",
        "last_action_date":     datetime.now(timezone.utc).isoformat(),
    }
    try:
        await asyncio.to_thread(insert_launch, row)
    except Exception as exc:
        raise HTTPException(status_code=500, detail=str(exc))

    return {"success": True, "kol_id": req.kol_id, "handle": kol.get("creator_display_name")}


# ---------------------------------------------------------------------------
# Report generator — Excel download split by SG / MY
# ---------------------------------------------------------------------------


@app.get("/api/reports/download")
async def download_report() -> StreamingResponse:
    """Pull launches table, split by market, and stream as Excel."""
    try:
        records = await asyncio.to_thread(read_launches)
    except Exception:
        records = []

    wb = openpyxl.Workbook()
    headers = [
        "Brand", "Product / Signal", "Market", "Retail Channel",
        "Distributor", "Festive Tag", "Pitch Status", "Creator", "Last Action",
    ]

    for sheet_label, market_key in [("SG Launches", "Singapore"), ("MY Launches", "Malaysia")]:
        ws = wb.create_sheet(title=sheet_label)
        ws.append(headers)
        for cell in ws[1]:
            cell.font = openpyxl.styles.Font(bold=True)

        for r in records:
            if r.get("market", "Singapore") != market_key:
                continue
            ws.append([
                r.get("brand_name", ""),
                r.get("product_name", ""),
                r.get("market", ""),
                r.get("retail_channel", ""),
                r.get("distributor_point", "") or _detect_distributor(r.get("brand_name", "")),
                r.get("festive_tag", "None") or _tag_festive(r.get("product_name", "")),
                r.get("pitch_status", "Raw"),
                r.get("creator_display_name", ""),
                r.get("last_action_date", ""),
            ])

    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"elitez_regional_intel_{datetime.now(timezone.utc).strftime('%Y%m%d')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# ---------------------------------------------------------------------------
# Internal scraper logic (sync — called via asyncio.to_thread)
# ---------------------------------------------------------------------------


def _run_scrape() -> dict:
    inserted = 0
    skipped = 0

    for feed_url in _RSS_FEEDS:
        try:
            feed = feedparser.parse(feed_url)
        except Exception:
            continue

        for entry in feed.entries:
            title   = getattr(entry, "title",   "")
            url     = getattr(entry, "link",    "")
            summary = getattr(entry, "summary", "")

            if not _is_fmcg_relevant(title, summary):
                skipped += 1
                continue

            combined = title + " " + summary
            row: dict = {
                "id":                   compute_url_hash(url),
                "brand_name":           _detect_brand(combined),
                "product_name":         title,
                "market":               _detect_market(combined),
                "retail_channel":       _detect_retail_channel(combined),
                "distributor_point":    _detect_distributor(combined),
                "festive_tag":          _tag_festive(combined),
                "product_photo_url":    None,
                "creator_display_name": None,
                "pitch_status":         "Raw",
                "last_action_date":     datetime.now(timezone.utc).isoformat(),
            }
            try:
                upsert_launch(row)
                inserted += 1
            except Exception:
                skipped += 1

    return {"inserted": inserted, "skipped": skipped}
